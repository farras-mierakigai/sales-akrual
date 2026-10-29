"""Tes analisa pelanggan berulang: identitas pembeli di parser (Shopee utuh, TikTok tersamar), API /api/repeat, parse ulang file lama.

Jalankan dari root repo:  python3 tests/test_repeat.py   (atau pytest). Pakai SQLite sementara.
"""
import os
import sys
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

if "db" in sys.modules:
    import db
    TMP = db.DATA
else:
    TMP = tempfile.mkdtemp(prefix="sa-repeat-")
    os.environ.update(SA_MODE="server", SA_DATA_DIR=TMP, SA_DB_URL="sqlite:///" + os.path.join(TMP, "t.db"),
                      SA_ADMIN_USER="farras", SA_ADMIN_PASS="rahasia-tim-2026", SA_SECRET="x" * 32,
                      SA_GOOGLE_CLIENT_ID="id", SA_GOOGLE_CLIENT_SECRET="secret", SA_BASE_URL="https://apps.test")

import pandas as pd                      # noqa: E402
import sqlalchemy as sa                  # noqa: E402
import db, auth, parsers, analysis, app as appmod   # noqa: E402
from fastapi.testclient import TestClient   # noqa: E402

auth.bootstrap_from_env()
TOKO = "TokoRepeat"


def farras():
    c = TestClient(appmod.app, base_url="https://apps.test")
    assert c.post("/api/login", json=dict(username="farras", password="rahasia-tim-2026")).status_code == 200
    return c


def shopee_xlsx(path, rows):
    """rows: (order_id, tanggal, status, username, penerima, telepon, produk, qty, subtotal)"""
    df = pd.DataFrame([dict({"No. Pesanan": o, "Waktu Pesanan Dibuat": d, "Status Pesanan": st, "Status Pembatalan/ Pengembalian": "",
                             "Nomor Referensi SKU": "SKU1", "Nama Produk": p, "Nama Variasi": "", "Jumlah": q, "Subtotal Pesanan": sub,
                             "Harga Awal": sub, "Metode Pembayaran": "COD", "Alasan Pembatalan": "", "Username (Pembeli)": u,
                             "Nama Penerima": r, "No. Telepon": t}) for o, d, st, u, r, t, p, q, sub in rows])
    df.to_excel(path, index=False)


def tiktok_csv(path, rows):
    """rows: (order_id, created 'dd/mm/yyyy HH:MM:SS', status, buyer_username, recipient, phone, product, qty, subtotal)"""
    cols = ["Order ID", "Order Status", "Created Time", "SKU ID", "Seller SKU", "Product Name", "Variation", "Quantity",
            "SKU Subtotal Before Discount", "SKU Seller Discount", "SKU Subtotal After Discount", "Payment Method", "Cancel Reason",
            "Cancelation/Return Type", "Purchase Channel", "Buyer Username", "Recipient", "Phone #"]
    df = pd.DataFrame([[o, st, d, "9001", "SKU1", p, "", q, sub, 0, sub, "COD", "", "", "", u, r, t] for o, d, st, u, r, t, p, q, sub in rows], columns=cols)
    df.to_csv(path, index=False)


def test_parser_keeps_buyer_identity():
    sp = os.path.join(TMP, "Order.all.test.xlsx")
    shopee_xlsx(sp, [("260101AAA", "2026-01-05 10:00", "Selesai", "budi_88", "B***i", "******12", "Kemeja", 1, 100000),
                     ("260102BBB", "2026-02-05 10:00", "Selesai", "", "S***i", "******34", "Kemeja", 1, 100000)])
    o = parsers.parse_shopee_orders(sp)
    assert list(o["buyer"]) == ["budi_88", "S***i"]
    assert o["buyer_key"][0] == "budi_88"
    assert o["buyer_key"][1] == "|s***i|******34"           # tanpa username → gabungan penerima + telepon tersamar
    tp = os.path.join(TMP, "Semua pesanan-test.csv")
    tiktok_csv(tp, [("579000000001", "05/01/2026 10:00:00", "Selesai", "D***i a***ini", "D****", "(+62)812*****60", "Rok", 1, 50000),
                    ("579000000002", "06/02/2026 10:00:00", "Selesai", "D***i a***ini", "D****", "(+62)812*****60", "Rok", 2, 100000),
                    ("579000000003", "07/02/2026 10:00:00", "Selesai", "D***i a***ini", "D****", "(+62)813*****99", "Rok", 1, 50000)])
    assert parsers.detect(tp) == "tiktok_orders"
    t = parsers.parse_tiktok_orders(tp)
    assert t["buyer_key"][0] == t["buyer_key"][1]            # nama + telepon tersamar sama → pelanggan sama
    assert t["buyer_key"][2] != t["buyer_key"][0]            # telepon beda → pelanggan beda walau pola nama sama


def test_repeat_analysis_rules():
    def ol(oid, created, buyer, ch=parsers.SHOPEE, bruto=100.0, batal=False, unpaid=False, retur=False):
        return dict(platform=parsers.SHOPEE if ch == parsers.SHOPEE else parsers.TIKTOK, channel=ch, order_id=oid,
                    created=pd.Timestamp(created), month=created[:7], status="x", is_batal=batal, is_unpaid=unpaid, is_retur=retur,
                    product="P", variation="", qty=1.0, bruto=bruto, buyer=buyer, buyer_key=buyer.lower())
    O = pd.DataFrame([
        ol("A1", "2026-01-03", "ani"), ol("A2", "2026-01-20", "ani"), ol("A3", "2026-03-02", "ani"),      # 3× (jarak 17, 41 hari)
        ol("B1", "2026-01-10", "budi"), ol("B2", "2026-02-10", "budi", batal=True),                      # batal tidak dihitung
        ol("C1", "2026-02-15", "citra"), ol("C2", "2026-02-15", "citra", unpaid=True),
        ol("D1", "2026-02-20", "ani", ch=parsers.TIKTOK), ol("D2", "2026-03-20", "ani", ch=parsers.TIKTOK),   # username sama di TikTok = pelanggan lain
        ol("E1", "2026-03-05", "eko", retur=True),
    ])
    R = analysis.repeat_customers(O)
    T = R["totals"]
    assert T["customers"] == 4 and T["orders"] == 7                                  # ani(S) budi citra ani(T); 7 order terjual
    assert T["repeat_customers"] == 2 and T["repeat_pct"] == 50.0                     # ani(S) 3×, ani(T) 2×
    assert T["repeat_orders"] == 3 and T["repeat_bruto"] == 300.0                     # A2 A3 D2
    assert T["new_customers"] == 4 and T["returning_customers"] == 0
    assert T["gap_median"] == 28.0 and T["gap_avg"] == 28.7                           # (17, 41, 28)
    by = {c["label"]: c for c in R["channels"]}
    assert by["Shopee"]["customers"] == 3 and by["TikTok Shop"]["customers"] == 1 and by["TikTok Shop"]["repeat_orders"] == 1
    m = {x["month"]: x for x in R["monthly"]}
    assert (m["2026-01"]["new"], m["2026-01"]["returning"]) == (2, 0)
    assert (m["2026-02"]["new"], m["2026-02"]["returning"]) == (2, 0)                 # citra & ani(T) baru; budi batal
    assert (m["2026-03"]["new"], m["2026-03"]["returning"]) == (0, 2)                 # ani(S) & ani(T) kembali
    dist = {d["bucket"]: d["customers"] for d in R["dist"]}
    assert dist["1"] == 2 and dist["2"] == 1 and dist["3"] == 1
    coh = {c["cohort"]: c for c in R["cohorts"]}
    assert coh["2026-01"]["size"] == 2 and coh["2026-01"]["pct"] == [100.0, 0.0, 50.0]   # ani kembali di +2 (A3), budi tidak
    assert coh["2026-02"]["pct"] == [100.0, 50.0]
    assert R["top"][0]["buyer"] == "ani" and R["top"][0]["orders_total"] == 3 and R["top"][0]["gap_avg"] == 29.0
    # filter bulan: urutan order tetap dari semua data → Maret: 2 pelanggan lama, 0 baru, semua order = order ulang
    R2 = analysis.repeat_customers(O, m_from="2026-03", m_to="2026-03")
    assert R2["totals"]["new_customers"] == 0 and R2["totals"]["returning_customers"] == 2 and R2["totals"]["repeat_orders_pct"] == 100.0
    # order tanpa identitas dilaporkan, bukan dianggap satu pelanggan
    O2 = pd.concat([O, pd.DataFrame([ol("Z1", "2026-03-09", "")])], ignore_index=True)
    R3 = analysis.repeat_customers(O2)
    assert R3["no_id"] == 1 and R3["totals"]["customers"] == 4


def test_repeat_habits_new_product_vs_new_variant():
    def ln(oid, created, buyer, product, variation="", ch=parsers.SHOPEE, bruto=100.0):
        return dict(platform=parsers.SHOPEE if ch == parsers.SHOPEE else parsers.TIKTOK, channel=ch, order_id=oid,
                    created=pd.Timestamp(created), month=created[:7], status="x", is_batal=False, is_unpaid=False, is_retur=False,
                    product=product, variation=variation, qty=1.0, bruto=bruto, buyer=buyer, buyer_key=buyer.lower())
    O = pd.DataFrame([
        ln("A1", "2026-01-03", "ani", "Kemeja", "Black"),
        ln("A2", "2026-01-20", "ani", "Kemeja", "Navy"),                       # warna baru, produk sama
        ln("A3", "2026-02-02", "ani", "Kemeja", "black "),                     # sama persis (beda kapital/spasi tetap sama)
        ln("A4", "2026-03-02", "ani", "Celana", "Black"), ln("A4", "2026-03-02", "ani", "Kemeja", "Navy"),   # campuran → produk baru
        ln("B1", "2026-01-10", "budi", "Kemeja", "Black"),
        ln("B2", "2026-02-10", "budi", "Kemeja", "Black"),                     # sama persis
        ln("C1", "2026-02-15", "citra", "Rok", ch=parsers.TIKTOK),
        ln("C2", "2026-03-15", "citra", "Rok", ch=parsers.TIKTOK),             # TikTok tanpa varian → sama persis
    ])
    R = analysis.repeat_customers(O)
    H = R["habits"]
    T = H["totals"]
    assert T["orders"] == 5 and T["new"] == 1 and T["variant"] == 1 and T["same"] == 3
    assert T["new_pct"] == 20.0 and T["new_bruto"] == 200.0
    by = {c["label"]: c for c in H["channels"]}
    assert by["Shopee"]["orders"] == 4 and by["TikTok Shop"]["same"] == 1
    seq = {s["bucket"]: s for s in H["seq"]}
    assert seq["ke-2"]["orders"] == 3 and seq["ke-2"]["variant"] == 1 and seq["ke-2"]["same"] == 2
    assert seq["ke-4–5"]["new"] == 1
    prods = {(p["channel"], p["product"]): p for p in H["products"]}
    assert prods[("Shopee", "Kemeja")]["customers"] == 2 and prods[("Shopee", "Kemeja")]["lines"] == 4     # A2 A3 A4(Kemeja) B2
    assert prods[("Shopee", "Kemeja")]["variant"] == 1 and prods[("Shopee", "Kemeja")]["same"] == 3
    assert H["new_products"][0]["product"] == "Celana" and H["new_products"][0]["customers"] == 1
    assert H["pairs"] == [dict(channel="Shopee", src="Kemeja", dst="Celana", customers=1)]
    top = {t["buyer"]: t for t in R["top"]}
    assert top["ani"]["habit"] == dict(new=1, variant=1, same=1)
    assert top["citra"]["habit"] == dict(new=0, variant=0, same=1)
    # periode: riwayat tetap dari semua data; Maret hanya A4 (produk baru) & C2 (sama)
    H2 = analysis.repeat_customers(O, m_from="2026-03", m_to="2026-03")["habits"]
    assert H2["totals"]["orders"] == 2 and H2["totals"]["new"] == 1 and H2["totals"]["same"] == 1
    assert H2["products"][0]["product"] in ("Kemeja", "Rok")


def test_product_launches_new_vs_old_buyers():
    def ln(oid, created, buyer, product, ch=parsers.SHOPEE, bruto=100.0):
        return dict(platform=parsers.SHOPEE, channel=ch, order_id=oid, created=pd.Timestamp(created), month=created[:7], status="x",
                    is_batal=False, is_unpaid=False, is_retur=False, product=product, variation="", qty=1.0, bruto=bruto,
                    buyer=buyer, buyer_key=buyer.lower())
    O = pd.DataFrame([
        ln("A1", "2026-01-03", "ani", "Kemeja"), ln("B1", "2026-01-10", "budi", "Kemeja"), ln("C1", "2026-02-10", "citra", "Kemeja"),
        # "Celana" rilis Maret (tidak ada penjualan Jan–Feb): ani & budi = pelanggan lama, dewi & eko = orang baru; eko beli lagi April
        ln("A2", "2026-03-02", "ani", "Celana"), ln("B2", "2026-03-05", "budi", "Celana"), ln("D1", "2026-03-06", "dewi", "Celana"),
        ln("E1", "2026-04-01", "eko", "Celana"), ln("E2", "2026-04-20", "eko", "Celana"),
        # "Rok" hanya 1 order di bulan pertamanya → di bawah ambang, bukan rilis; "Kemeja" ada sejak bulan pertama data → bukan rilis
        ln("F1", "2026-04-03", "fani", "Rok"),
    ])
    LA = analysis.repeat_customers(O)["habits"]["launches"]
    assert [r["product"] for r in LA["items"]] == ["Celana"] and LA["items"][0]["month"] == "2026-03"
    r = LA["items"][0]
    assert r["buyers"] == 4 and r["new"] == 2 and r["old"] == 2 and r["new_pct"] == 50.0 and r["orders"] == 5
    assert LA["totals"]["launches"] == 1 and LA["totals"]["buyers"] == 4 and LA["channels"][0]["label"] == "Shopee"
    # periode April saja: rilis Maret tidak masuk periode → kosong
    assert analysis.repeat_customers(O, m_from="2026-04", m_to="2026-04")["habits"]["launches"]["items"] == []
    # periode Maret saja: jendela pembeli dibatasi Maret → 3 pembeli (eko belum)
    r = analysis.repeat_customers(O, m_from="2026-03", m_to="2026-03")["habits"]["launches"]["items"][0]
    assert r["buyers"] == 3 and r["new"] == 1


def test_api_repeat_and_reparse_old_uploads():
    c = farras()
    assert c.post("/api/toko", json=dict(name=TOKO)).status_code == 200
    d = db.archive_dir(TOKO); os.makedirs(d, exist_ok=True)
    sp = os.path.join(d, "Order.all.20260101_20260228.xlsx")
    shopee_xlsx(sp, [("260101AAA", "2026-01-05 10:00", "Selesai", "budi_88", "B***i", "******12", "Kemeja", 1, 100000),
                     ("260201BBB", "2026-02-05 10:00", "Selesai", "budi_88", "B***i", "******12", "Celana", 1, 150000),
                     ("260202CCC", "2026-02-06 10:00", "Selesai", "sari", "S***i", "******34", "Kemeja", 1, 100000)])
    db.ingest_file(TOKO, sp, uploaded_by="farras")
    r = c.get("/api/repeat", params=dict(toko=TOKO)).json()
    assert r["totals"]["customers"] == 2 and r["totals"]["repeat_customers"] == 1 and r["reparsed"] == 0
    assert r["top"][0]["buyer"] == "budi_88" and r["top"][0]["product"] in ("Kemeja", "Celana")
    # simulasi unggahan dari parser lama: kolom pembeli kosong + parser_version lama → diparse ulang otomatis, pengunggah dipertahankan
    with db.engine().begin() as cx:
        cx.execute(sa.text("UPDATE order_line SET buyer=NULL, buyer_key=NULL"))
        cx.execute(sa.text("UPDATE upload SET parser_version='10'"))
    assert db.uploads_needing_buyer(TOKO) == [os.path.basename(sp)]
    r = c.get("/api/repeat", params=dict(toko=TOKO)).json()
    assert r["reparsed"] == 1 and r["totals"]["customers"] == 2
    assert db.uploads_needing_buyer(TOKO) == []
    files = c.get("/api/files", params=dict(toko=TOKO)).json()
    assert files[0]["uploaded_by"] == "farras"
    x = c.get("/api/repeat.xlsx", params=dict(toko=TOKO))
    assert x.status_code == 200 and x.headers["content-type"].startswith("application/vnd.openxmlformats")
    from openpyxl import load_workbook
    import io
    assert "Kebiasaan order ulang" in load_workbook(io.BytesIO(x.content)).sheetnames
    # workspace lain tidak melihat toko ini
    assert c.get("/api/repeat", params=dict(toko="TokoLain")).status_code == 404


if __name__ == "__main__":
    for name, fn in list(globals().items()):
        if name.startswith("test_") and callable(fn):
            fn()
            print("ok", name)
    print("semua tes lulus")
