"""Sistem audit: kelengkapan data, rekonsiliasi vs total resmi platform, temuan anomali per order, jejak order."""
import calendar
import datetime as dt
import pandas as pd
import analysis
from parsers import SHOPEE, TIKTOK, TOKOPEDIA

SHOPEE_FEE_KEYS = [k for k, _ in analysis.FEE_LINES[SHOPEE]]
MAX_ROWS = 300


def _rows(df, cols, n=MAX_ROWS):
    if not len(df):
        return []
    d = df[cols].head(n).copy()
    for c in d.columns:
        if str(d[c].dtype).startswith("datetime"):
            d[c] = d[c].dt.strftime("%Y-%m-%d").fillna("")
    return d.to_dict("records")


def _filter_months(df, col, m_from, m_to):
    if not len(df):
        return df
    if m_from:
        df = df[df[col] >= m_from]
    if m_to:
        df = df[df[col] <= m_to]
    return df


# ---------------------------------------------------------------- 1. kelengkapan data

def coverage(O, I, m_from=None, m_to=None, today=None):
    today = today or dt.date.today()
    out = []
    if not len(O) and not len(I):
        return out
    I = I.copy() if len(I) else I
    if len(I):
        om = O.drop_duplicates("order_id").set_index("order_id")["month"] if len(O) else pd.Series(dtype=str)
        I["month"] = I["order_id"].map(om).fillna(I["created"].dt.strftime("%Y-%m"))
    plats = sorted(set(O["platform"]) | (set(I["platform"]) if len(I) else set()))
    for p in plats:
        o = O[O["platform"] == p] if len(O) else O
        i = I[I["platform"] == p] if len(I) else I
        months = sorted((set(o["month"].dropna()) if len(o) else set()) | (set(i["month"].dropna()) if len(i) else set()))
        if m_from:
            months = [m for m in months if m >= m_from]
        if m_to:
            months = [m for m in months if m <= m_to]
        for m in months:
            y, mm = int(m[:4]), int(m[5:7])
            ndays = calendar.monthrange(y, mm)[1]
            last_day = min(ndays, today.day) if (y, mm) == (today.year, today.month) else ndays
            if (y, mm) > (today.year, today.month):
                last_day = 0
            om_ = o[o["month"] == m] if len(o) else o
            im_ = i[i["month"] == m] if len(i) else i
            days_with = set(om_["created"].dt.day.dropna().astype(int)) if len(om_) else set()
            missing = [d for d in range(1, last_day + 1) if d not in days_with]
            live = om_[~om_["is_unpaid"]] if len(om_) else om_
            sold = live[(~live["is_batal"]) & (~live["is_retur"])] if len(live) else live
            settled_ids = set(im_["order_id"]) if len(im_) else set()
            uns = sold[~sold["order_id"].isin(settled_ids)] if len(sold) else sold
            bruto = float(sold["bruto"].sum()) if len(sold) else 0.0
            uns_b = float(uns["bruto"].sum()) if len(uns) else 0.0
            r = dict(platform=p, month=m, order_rows=int(len(om_)), orders=int(live["order_id"].nunique()) if len(live) else 0,
                     days_total=last_day, days_with=len(days_with), missing_days=missing,
                     income_orders=int(im_["order_id"].nunique()) if len(im_) else 0,
                     settle_pct=round(100.0 * (1 - uns_b / bruto), 1) if bruto else None,
                     income_tanpa_pesanan=int((~im_["order_id"].isin(set(om_["order_id"]))).sum()) if len(im_) and len(om_) else (int(im_["order_id"].nunique()) if len(im_) else 0),
                     files_orders=sorted(set(om_["file"])) if len(om_) else [],
                     files_income=sorted(set(im_["file"])) if len(im_) else [])
            if not len(om_):
                r["status"] = "bad"; r["note"] = "Tidak ada file pesanan bulan ini (hanya laporan penghasilan)."
            elif len(missing) >= 3:
                r["status"] = "bad"; r["note"] = f"{len(missing)} hari tanpa satu pun order — kemungkinan export pesanan tidak menutup seluruh bulan."
            elif missing:
                r["status"] = "warn"; r["note"] = f"{len(missing)} hari tanpa order (tgl {', '.join(map(str, missing))})."
            elif not len(im_):
                r["status"] = "warn"; r["note"] = "Belum ada laporan penghasilan untuk bulan ini — potongan/neto belum terhitung."
            elif r["settle_pct"] is not None and r["settle_pct"] < 95:
                r["status"] = "warn"; r["note"] = f"Baru {r['settle_pct']}% bruto yang muncul di laporan penghasilan."
            else:
                r["status"] = "ok"; r["note"] = "Lengkap."
            out.append(r)
    return out


# ---------------------------------------------------------------- 2. rekonsiliasi vs total resmi per file

LABELS = {
    "harga_produk": "Harga produk / subtotal setelah diskon penjual", "refund": "Refund ke pembeli",
    "biaya_admin": "Biaya administrasi", "biaya_proses": "Biaya proses pesanan", "biaya_xtra": "Biaya gratis ongkir XTRA",
    "komisi_ams": "Komisi AMS", "pph22": "PPh 22", "neto": "Total yang dilepas (neto)",
    "settlement_total": "Total settlement (order + penyesuaian)", "total_fees": "Total biaya",
    "komisi_platform": "Komisi platform", "komisi_dinamis": "Komisi dinamis", "ongkir_net": "Ongkir (net)",
}


def reconcile_files(raw, files):
    out = []
    RI, RA = raw.get("income"), raw.get("adj")
    for f in files:
        summ = f.get("summary") or {}
        if not summ:
            continue
        ri = RI[RI["file"] == f["file"]] if len(RI) else RI
        ra = RA[RA["file"] == f["file"]] if len(RA) else RA
        items = []
        summ = dict(summ)
        if "diskon_produk" in summ:          # format lama: ringkasan pakai harga sebelum diskon → samakan basis (setelah diskon penjual)
            summ["harga_produk"] = summ.get("harga_produk", 0) + summ.pop("diskon_produk")
        for k, official in summ.items():
            if k == "settlement_total":
                parsed = float(ri["neto"].sum() if len(ri) else 0) + float(ra["amount"].sum() if len(ra) else 0)
            elif k in ri.columns if len(ri) else False:
                parsed = float(ri[k].sum())
            else:
                continue
            diff = parsed - official
            items.append(dict(key=k, label=LABELS.get(k, k), official=official, parsed=parsed, diff=diff, ok=abs(diff) <= 1))
        out.append(dict(file=f["file"], label=f["label"], seller=f.get("seller", ""), items=items,
                        ok=all(x["ok"] for x in items), n_items=len(items), n_bad=sum(1 for x in items if not x["ok"])))
    return out


# ---------------------------------------------------------------- 3. temuan per order

def findings(O, I, raw, m_from=None, m_to=None, today=None):
    today = pd.Timestamp(today or dt.date.today())
    F = []
    O = _filter_months(O, "month", m_from, m_to) if len(O) else O
    sold = O[(~O["is_unpaid"]) & (~O["is_batal"]) & (~O["is_retur"])] if len(O) else O
    status_by_order = O.drop_duplicates("order_id").set_index("order_id")["status"] if len(O) else pd.Series(dtype=str)
    batal_by_order = O.groupby("order_id")["is_batal"].all() if len(O) else pd.Series(dtype=bool)

    def add(fid, title, sev, df, cols, desc, amount_col=None, note=None, extra=None):
        F.append(dict(id=fid, title=title, severity=sev if len(df) else "ok", count=int(len(df)),
                      amount=float(df[amount_col].sum()) if (amount_col and len(df)) else None,
                      desc=desc, note=note, rows=_rows(df, cols), cols=cols, extra=extra or {}))

    # income per order (gabung baris)
    if len(I):
        Ig = I.groupby(["platform", "order_id"]).agg(created=("created", "min"), settled=("settled", "max"),
                                                      harga_produk=("harga_produk", "sum"), refund=("refund", "sum"),
                                                      neto=("neto", "sum"), file=("file", "last")).reset_index()
        fee = pd.Series(0.0, index=I.index)
        for k in SHOPEE_FEE_KEYS:
            if k in I:
                fee = fee + I[k].fillna(0)
        if "total_fees" in I:
            fee = fee.where(I["platform"] != TIKTOK, I["total_fees"].fillna(0))
        Ig["potongan"] = I.assign(_f=fee).groupby(["platform", "order_id"])["_f"].sum().values
        Ig["status"] = Ig["order_id"].map(status_by_order).fillna("(tidak ada di file pesanan)")
        Ig["month"] = Ig["order_id"].map(O.drop_duplicates("order_id").set_index("order_id")["month"]) if len(O) else None
        Ig["month"] = Ig["month"].fillna(Ig["created"].dt.strftime("%Y-%m"))
        Ig = _filter_months(Ig, "month", m_from, m_to)
    else:
        Ig = pd.DataFrame(columns=["platform", "order_id", "created", "settled", "harga_produk", "refund", "neto", "file", "potongan", "status", "month"])

    # F1 ada di laporan penghasilan tapi tidak ada di file pesanan
    d = Ig[~Ig["order_id"].isin(set(O["order_id"]))] if len(O) else Ig
    add("income_tanpa_pesanan", "Order ada di laporan penghasilan tapi tidak ada di file pesanan", "bad", d,
        ["platform", "order_id", "created", "month", "harga_produk", "neto", "file"],
        "Order ini ikut dihitung di potongan/neto bulan tersebut, tapi bruto & statusnya tidak diketahui. Biasanya file pesanan periode itu belum diunggah.",
        amount_col="harga_produk")

    # F2 terjual tapi belum muncul di laporan penghasilan > 30 hari
    if len(sold):
        so = sold.groupby(["platform", "order_id"]).agg(created=("created", "min"), status=("status", "first"), bruto=("bruto", "sum"), product=("product", "first")).reset_index()
        so = so[~so["order_id"].isin(set(I["order_id"]) if len(I) else set())]
        so = so[so["created"] < today - pd.Timedelta(days=30)]
        add("belum_settle_lama", "Terjual > 30 hari lalu tapi belum ada di laporan penghasilan", "warn", so.sort_values("created"),
            ["platform", "order_id", "created", "status", "bruto", "product"],
            "Bruto sudah diakui, tapi potongan/refund/neto order ini belum ada. Upload laporan penghasilan periode berikutnya, atau cek apakah order ini retur/gagal kirim.",
            amount_col="bruto")
    else:
        add("belum_settle_lama", "Terjual > 30 hari lalu tapi belum ada di laporan penghasilan", "warn", pd.DataFrame(), [], "")

    # F3 refund penuh tapi status terjual
    d = Ig[(Ig["harga_produk"] > 0) & (Ig["refund"] <= -0.99 * Ig["harga_produk"])]
    d = d[~d["order_id"].map(batal_by_order).fillna(False).astype(bool)]
    d = d[d["order_id"].isin(set(O["order_id"]))] if len(O) else d
    add("refund_penuh_terjual", "Retur penuh: dana dikembalikan penuh, status pesanan bukan batal → dikeluarkan dari bruto", "info", d,
        ["platform", "order_id", "created", "status", "harga_produk", "refund", "neto"],
        "Aturan (Farras, 2026-10-08): order ini diperlakukan sebagai retur — TIDAK masuk bruto dan tidak dihitung sebagai pesanan terjual, walau status di file pesanan non-batal. Muncul di laporan sebagai baris 'Retur penuh'.",
        amount_col="harga_produk")

    # F4 neto negatif
    d = Ig[Ig["neto"] < 0]
    add("neto_negatif", "Dana diterima negatif per order", "warn", d.sort_values("neto"),
        ["platform", "order_id", "created", "status", "harga_produk", "refund", "potongan", "neto"],
        "Penjual nombok di order ini (biasanya retur + ongkir balik, atau penalti). Cek alasannya di seller center.",
        amount_col="neto")

    # F5 potongan tak wajar
    d = Ig[Ig["harga_produk"] > 0].copy()
    if len(d):
        d["potongan_pct"] = (-d["potongan"] / d["harga_produk"] * 100).round(1)
        med = d.groupby("platform")["potongan_pct"].median().round(1).to_dict()
        bad = d[(d["potongan_pct"] > 50) | (d["potongan_pct"] < 0)].sort_values("potongan_pct", ascending=False)
        add("potongan_tak_wajar", "Potongan marketplace > 50% atau negatif per order", "warn", bad,
            ["platform", "order_id", "created", "status", "harga_produk", "potongan", "potongan_pct", "neto"],
            "Median potongan per platform: " + ", ".join(f"{k} {v}%" for k, v in med.items()) + ". Order di luar batas wajar perlu dicek (komisi afiliasi/iklan besar, atau salah pasang order).",
            extra=dict(median=med))
    else:
        add("potongan_tak_wajar", "Potongan marketplace > 50% atau negatif per order", "warn", pd.DataFrame(), [], "")

    # F6 batal tapi ada dana masuk
    if len(Ig) and len(O):
        d = Ig[Ig["order_id"].map(batal_by_order).fillna(False).astype(bool) & (Ig["neto"] > 0)]
        add("batal_ada_dana", "Status batal di file pesanan, tapi ada dana masuk di laporan penghasilan", "bad", d,
            ["platform", "order_id", "created", "status", "harga_produk", "neto", "file"],
            "Konflik status: kemungkinan file pesanan lama (status belum update) atau batal sebagian. Order ini TIDAK masuk bruto tapi netonya terhitung.",
            amount_col="neto")

    # F7 status berubah antar file export
    RO = raw.get("orders")
    if len(RO):
        ro = RO.sort_values("file_mtime")
        g = ro.groupby("line_key")
        multi = g["status_raw"].nunique()
        keys = multi[multi > 1].index
        if len(keys):
            sub = ro[ro["line_key"].isin(keys)]
            first = sub.groupby("line_key").first()
            last = sub.groupby("line_key").last()
            ch = pd.DataFrame({"platform": last["platform"], "order_id": last["order_id"], "product": last["product"],
                               "dari": first["status"], "ke": last["status"], "file_lama": first["file"], "file_baru": last["file"]}).reset_index(drop=True)
            trans = ch.groupby(["dari", "ke"]).size().sort_values(ascending=False).reset_index(name="n").to_dict("records")
            ke_batal = ch[ch["ke"].str.lower().str.contains("batal")]
            add("status_berubah", "Status berubah antar file export (dipakai: file terbaru)", "info", ch,
                ["platform", "order_id", "dari", "ke", "file_lama", "file_baru", "product"],
                f"{len(ch)} baris berubah status; {len(ke_batal)} di antaranya berubah menjadi batal. Ini normal kalau export lama diambil sebelum order selesai — pastikan export terbaru menutup semua bulan.",
                extra=dict(transitions=trans))
        else:
            add("status_berubah", "Status berubah antar file export", "info", pd.DataFrame(), [], "Tidak ada perubahan status antar file.")

    # F8 bruto file pesanan vs harga produk laporan penghasilan
    rc = analysis.reconcile(O, I)
    mm = pd.DataFrame(rc["mismatch"])
    add("bruto_vs_income", "Bruto per order (file pesanan) ≠ harga produk (laporan penghasilan)", "bad", mm,
        ["order_id", "orders", "income", "diff"] if len(mm) else [],
        f"{rc['exact']} dari {rc['matched']} order cocok persis. Selisih berarti definisi bruto meleset atau ada order sebagian.",
        extra=dict(matched=rc["matched"], exact=rc["exact"]))

    # F9 terjual tapi nilai 0
    d = sold[sold["bruto"] <= 0] if len(sold) else sold
    add("terjual_nol", "Baris terjual dengan nilai Rp 0", "info", d,
        ["platform", "order_id", "created", "status", "product", "qty", "bruto"],
        "Biasanya hadiah/bonus/voucher 100%. Tidak menambah bruto, tapi menambah qty.")

    # F10 tanggal kosong
    d = O[O["created"].isna()] if len(O) else O
    add("tanggal_kosong", "Baris pesanan tanpa tanggal order", "bad", d,
        ["platform", "order_id", "status", "product", "bruto", "file"],
        "Baris ini tidak bisa masuk bulan mana pun → tidak terhitung.")

    # F11 duplikat order ID antar platform (tidak mungkin, jaga-jaga)
    if len(O):
        dup = O.drop_duplicates(["platform", "order_id"]).groupby("order_id")["platform"].nunique()
        dup = dup[dup > 1]
        d = O[O["order_id"].isin(dup.index)].drop_duplicates(["platform", "order_id"])
        add("order_id_lintas_platform", "Order ID sama muncul di dua platform", "bad", d,
            ["platform", "order_id", "created", "status", "bruto", "file"], "Kemungkinan file salah masuk toko/platform.")

    summary = dict(bad=sum(1 for f in F if f["severity"] == "bad"), warn=sum(1 for f in F if f["severity"] == "warn"),
                   info=sum(1 for f in F if f["severity"] == "info"), ok=sum(1 for f in F if f["severity"] == "ok"))
    return dict(findings=F, summary=summary)


# ---------------------------------------------------------------- 4. jejak order

def trace(O, I, A, raw, order_id):
    oid = str(order_id).strip()
    RO, RI = raw.get("orders"), raw.get("income")
    res = dict(order_id=oid, found=False)
    if len(RO):
        r = RO[RO["order_id"] == oid].sort_values(["file_mtime", "line_key"])
        res["raw_orders"] = _rows(r, ["file", "status_raw", "product", "variation", "qty", "bruto", "created", "payment", "cancel_reason"])
    if len(O):
        o = O[O["order_id"] == oid]
        if len(o):
            res["found"] = True
            live = o[~o["is_unpaid"]]
            res["final"] = dict(platform=o["platform"].iloc[0], channel=o["channel"].iloc[0], month=o["month"].iloc[0],
                                status=o["status"].iloc[0], is_batal=bool(o["is_batal"].all()), is_unpaid=bool(o["is_unpaid"].all()),
                                is_retur=bool(o["is_retur"].all()),
                                bruto=float(live[(~live["is_batal"]) & (~live["is_retur"])]["bruto"].sum()),
                                qty=float(live[(~live["is_batal"]) & (~live["is_retur"])]["qty"].sum()),
                                nilai_retur=float(live[live["is_retur"]]["bruto"].sum()),
                                nilai_batal=float(live[live["is_batal"]]["bruto"].sum()), lines=_rows(o, ["product", "variation", "qty", "bruto", "status", "file"]))
    if len(RI):
        i = RI[RI["order_id"] == oid].sort_values("file_mtime")
        cols = [c for c in ["file", "created", "settled", "harga_produk", "refund", "ongkir_net", "voucher_penjual", "biaya_admin", "biaya_proses", "biaya_xtra", "biaya_layanan", "komisi_ams", "promosi_lain", "komisi_platform", "komisi_dinamis", "afiliasi", "biaya_lainnya", "pph22", "neto"] if c in i.columns]
        res["raw_income"] = _rows(i, cols)
        if len(i):
            res["found"] = True
    if len(I):
        i = I[I["order_id"] == oid]
        if len(i):
            res["income_final"] = dict(neto=float(i["neto"].sum()), refund=float(i["refund"].sum()), harga_produk=float(i["harga_produk"].sum()), n=int(len(i)))
    if len(A):
        a = A[(A["id"] == oid) | (A["related"] == oid)]
        res["adjustments"] = _rows(a, ["type", "settled", "amount", "file"])
    return res


# ---------------------------------------------------------------- 5. export audit excel

def export_audit_xlsx(toko, cov, rec, fin):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Ringkasan"
    ws.append([f"AUDIT DATA PENJUALAN — {toko}"]); ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Dibuat {dt.date.today().isoformat()}"]); ws.append([])
    ws.append(["Temuan", "Tingkat", "Jumlah", "Nilai (Rp)", "Keterangan"])
    for c in ws[4]: c.font = Font(bold=True)
    for f in fin["findings"]:
        ws.append([f["title"], f["severity"], f["count"], f["amount"], f["desc"]])
    ws.column_dimensions["A"].width = 60; ws.column_dimensions["E"].width = 90

    ws = wb.create_sheet("Kelengkapan")
    ws.append(["Platform", "Bulan", "Order", "Baris", "Hari ada order", "Hari dalam bulan", "Hari kosong", "Order di lap. penghasilan", "Settle %", "Status", "Catatan"])
    for c in ws[1]: c.font = Font(bold=True)
    for r in cov:
        ws.append([r["platform"], r["month"], r["orders"], r["order_rows"], r["days_with"], r["days_total"], ", ".join(map(str, r["missing_days"])), r["income_orders"], r["settle_pct"], r["status"], r["note"]])

    ws = wb.create_sheet("Rekonsiliasi")
    ws.append(["File", "Jenis", "Penjual", "Item", "Resmi (ringkasan platform)", "Hasil parse", "Selisih", "OK"])
    for c in ws[1]: c.font = Font(bold=True)
    for f in rec:
        for it in f["items"]:
            ws.append([f["file"], f["label"], f["seller"], it["label"], it["official"], it["parsed"], it["diff"], "ya" if it["ok"] else "TIDAK"])

    for f in fin["findings"]:
        if not f["rows"]:
            continue
        ws = wb.create_sheet(f["id"][:31])
        ws.append([f["title"]]); ws["A1"].font = Font(bold=True)
        ws.append([f["desc"]]); ws.append([])
        ws.append(f["cols"])
        for c in ws[4]: c.font = Font(bold=True)
        for r in f["rows"]:
            ws.append([r.get(c) for c in f["cols"]])
    return wb
