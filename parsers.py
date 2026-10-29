"""Parser export marketplace → DataFrame kanonis.

Format yang dikenali (deteksi dari isi file, bukan nama):
  - Shopee  Order.all*.xlsx            → orders  (1 baris = 1 produk/variasi)
  - Shopee  Income*.xlsx (Penghasilan) → income  (1 baris = 1 order, dana dilepas)
  - TikTok  "Semua pesanan*.csv"       → orders  (1 baris = 1 SKU; kolom Purchase Channel = TikTok/Tokopedia)
  - TikTok  income_*.xlsx (EN atau ID) → income  (1 baris = 1 order settle) + adjustment non-order
"""
import os
import re
import pickle
import hashlib
import warnings
import pandas as pd
import openpyxl

warnings.filterwarnings("ignore")

SHOPEE = "Shopee"
TIKTOK = "TikTok"
TOKOPEDIA = "Tokopedia"

# ---------------------------------------------------------------- util

_THOUSANDS = re.compile(r"^-?\d{1,3}(\.\d{3})+$")


def to_num(x):
    """Angka Indonesia ('1.165.080') maupun polos ('115200') → float."""
    if x is None:
        return 0.0
    if isinstance(x, (int, float)):
        return float(x)
    s = str(x).strip().replace(" ", "")
    if s in ("", "-", "/", "None", "nan"):
        return 0.0
    if _THOUSANDS.match(s):
        s = s.replace(".", "")
    s = s.replace(",", "")
    try:
        return float(s)
    except ValueError:
        return 0.0


def _s(x):
    return "" if x is None else str(x).strip()


def _sheet_rows(ws, n):
    return [list(r) for r in ws.iter_rows(min_row=1, max_row=n, values_only=True)]


def _find_header_row(rows, must_have):
    for i, r in enumerate(rows):
        cells = {_s(c) for c in r}
        if all(m in cells for m in must_have):
            return i
    return None


# ---------------------------------------------------------------- deteksi

def detect(path):
    """→ 'shopee_orders' | 'shopee_income' | 'tiktok_orders' | 'tiktok_income' | 'tiktok_unsettled' | None"""
    low = path.lower()
    if low.endswith(".csv"):
        with open(path, "rb") as fh:
            raw = fh.read(8192)
        if raw and raw.count(b"\x00") > len(raw) * 0.5:
            return "corrupt"
        head = raw.replace(b"\x00", b"").decode("utf-8-sig", "ignore").split("\n", 1)[0]
        if "Order ID" in head and "Order Status" in head and "SKU Subtotal After Discount" in head:
            return "tiktok_orders"
        return None
    if not low.endswith(".xlsx"):
        return None
    wb = openpyxl.load_workbook(path, read_only=False, data_only=True)
    try:
        names = wb.sheetnames
        if any(n.startswith("Penghasilan") for n in names) and "Summary" in names:
            return "shopee_income"
        if "Income" in names and "Summary" in names:
            return "shopee_income"
        ws = wb.worksheets[0]
        rows = _sheet_rows(ws, 8)
        for r in rows:
            cells = {_s(c) for c in r}
            if "No. Pesanan" in cells and "Status Pesanan" in cells:
                return "shopee_orders"
            if "Order ID" in cells and "Order Status" in cells and "SKU Subtotal After Discount" in cells:
                return "tiktok_orders"
            if ("Order/adjustment ID" in cells or "ID Pesanan/Penyesuaian" in cells) and (
                "Total settlement amount" in cells or "Jumlah penyelesaian pembayaran" in cells
            ):
                return "tiktok_income"
            if "ID pesanan terkait" in cells and any(c.startswith("Estimasi") for c in cells):
                return "tiktok_unsettled"
        return None
    finally:
        wb.close()


# ---------------------------------------------------------------- Shopee orders

def _norm_shopee_status(s):
    if s.startswith("Pesanan diterima"):
        return "Pesanan Diterima (window retur)"
    return s or "(kosong)"


def parse_shopee_orders(path):
    df = pd.read_excel(path, dtype=str)
    df.columns = [str(c).strip() for c in df.columns]
    df = df.fillna("")
    out = pd.DataFrame()
    out["platform"] = [SHOPEE] * len(df)
    out["channel"] = SHOPEE
    out["order_id"] = df["No. Pesanan"].str.strip()
    out["created"] = pd.to_datetime(df["Waktu Pesanan Dibuat"], errors="coerce")
    out["status_raw"] = df["Status Pesanan"].str.strip()
    out["status"] = out["status_raw"].map(_norm_shopee_status)
    out["is_batal"] = out["status_raw"].eq("Batal")
    out["is_unpaid"] = out["status_raw"].eq("Belum Bayar")
    out["is_return"] = df["Status Pembatalan/ Pengembalian"].str.strip().ne("")
    out["sku"] = df["Nomor Referensi SKU"].str.strip()
    out["product"] = df["Nama Produk"].str.strip()
    out["variation"] = df["Nama Variasi"].str.strip()
    out["qty"] = df["Jumlah"].map(to_num)
    if "Subtotal Pesanan" in df.columns:
        out["bruto"] = df["Subtotal Pesanan"].map(to_num)      # aturan tim: Subtotal Pesanan
    elif "Total Harga Produk" in df.columns:                     # format export lama (≤ Apr 2026)
        out["bruto"] = df["Total Harga Produk"].map(to_num)
    else:
        out["bruto"] = df["Harga Setelah Diskon"].map(to_num) * out["qty"]
    out["harga_awal"] = df["Harga Awal"].map(to_num) * out["qty"]
    out["payment"] = df["Metode Pembayaran"].str.strip()
    out["cancel_reason"] = df["Alasan Pembatalan"].str.strip()
    # identitas pembeli (buat analisa pelanggan berulang): Shopee memberi username utuh; nama penerima & telepon disamarkan
    uname = df["Username (Pembeli)"].str.strip() if "Username (Pembeli)" in df.columns else pd.Series([""] * len(df))
    rcpt = df["Nama Penerima"].str.strip() if "Nama Penerima" in df.columns else pd.Series([""] * len(df))
    phone = df["No. Telepon"].str.strip() if "No. Telepon" in df.columns else pd.Series([""] * len(df))
    out["buyer"] = uname.where(uname != "", rcpt)
    out["buyer_key"] = _buyer_key(uname, rcpt, phone)
    out["line_key"] = out["order_id"] + "|" + out["sku"] + "|" + out["variation"] + "|" + out["product"]
    return out


# akun agregat milik marketplace, bukan pelanggan: Shopee SIP (pembeli luar negeri lewat Shopee International Platform)
AGGREGATE_BUYERS = {"sipbuyerid"}


def _buyer_key(uname, rcpt, phone):
    """Kunci pelanggan: username kalau ada; kalau username disamarkan/kosong, gabungan nama penerima + telepon (tersamar) agar tetap bisa dicocokkan."""
    uname = uname.fillna("").astype(str).str.strip().str.lower()
    rcpt = rcpt.fillna("").astype(str).str.strip().str.lower()
    phone = phone.fillna("").astype(str).str.strip()
    masked = uname.eq("") | uname.str.contains(r"\*")
    alt = (uname + "|" + rcpt + "|" + phone).where(masked, uname)
    alt = alt.where((uname != "") | (rcpt != "") | (phone != ""), "")
    return alt.where(~uname.isin(AGGREGATE_BUYERS), "")


# ---------------------------------------------------------------- Shopee income

def parse_shopee_income(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    seller = ""
    try:
        for r in wb["Summary"].iter_rows(min_row=1, max_row=12, values_only=True):
            if _s(r[0]) == "Username (Penjual)":
                seller = _s(r[1])
    except Exception:
        pass
    frames = []
    if "Income" in wb.sheetnames and not any(n.startswith("Penghasilan") for n in wb.sheetnames):
        frames.append(_parse_shopee_income_old(wb["Income"]))
    for name in wb.sheetnames:
        if not name.startswith("Penghasilan"):
            continue
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        hi = _find_header_row(rows[:6], ["No. Pesanan", "Lihat berdasarkan"])
        if hi is None:
            continue
        header = [_s(c) for c in rows[hi]]
        group_row = [_s(c) for c in rows[hi - 1]] if hi >= 1 else [""] * len(header)
        # forward-fill nama grup (baris di atas header)
        groups, cur = [], ""
        for g in group_row:
            if g:
                cur = g
            groups.append(cur)
        idx = {h: i for i, h in enumerate(header)}
        i_view = idx["Lihat berdasarkan"]
        i_start = idx["Harga Produk"]
        i_end = idx["Username (Pembeli)"]
        i_total = idx.get("Total Penghasilan")

        def cols_in_group(gname):
            return [i for i in range(i_start, i_end) if groups[i] == gname and header[i] != "Total Penghasilan"]

        g_ongkir = cols_in_group("Subtotal Ongkos Kirim")
        g_voucher = cols_in_group("Voucher dan Potongan Harga")
        g_layanan = cols_in_group("Biaya Layanan")
        g_promosi = cols_in_group("Biaya Promosi")
        g_lainnya = cols_in_group("Biaya Lainnya")
        g_pajak = cols_in_group("Pajak")
        c_xtra = [i for i in range(i_start, i_end) if header[i].startswith("Biaya Gratis Ongkir XTRA")]
        c_admin = idx.get("Biaya Administrasi")
        c_proses = idx.get("Biaya Proses Pesanan")
        c_ams = idx.get("Biaya Komisi AMS")
        c_refund = idx.get("Jumlah Pengembalian Dana ke Pembeli")
        c_created = idx.get("Waktu Pesanan Dibuat")
        c_settled = idx.get("Tanggal Dana Dilepaskan")
        c_pengajuan = idx.get("No. Pengajuan")

        recs = []
        for r in rows[hi + 1:]:
            if _s(r[i_view]) != "Order":
                continue
            v = lambda i: to_num(r[i]) if i is not None else 0.0
            vs = lambda ii: sum(to_num(r[i]) for i in ii)
            neto = v(i_total) if i_total is not None else sum(to_num(r[i]) for i in range(i_start, i_end))
            rec = dict(
                platform=SHOPEE, channel=SHOPEE,
                order_id=_s(r[idx["No. Pesanan"]]),
                created=_s(r[c_created]) if c_created is not None else "",
                settled=_s(r[c_settled]) if c_settled is not None else "",
                neto=neto,
                harga_produk=v(i_start),
                refund=v(c_refund),
                ongkir_net=vs(g_ongkir),
                voucher_penjual=vs(g_voucher),
                biaya_admin=v(c_admin),
                biaya_proses=v(c_proses),
                biaya_xtra=vs(c_xtra),
                biaya_layanan=vs(g_layanan),
                komisi_ams=v(c_ams),
                promosi_lain=vs(g_promosi) - v(c_ams),
                biaya_lainnya=vs(g_lainnya),
                pph22=vs(g_pajak),
                pengajuan=_s(r[c_pengajuan]) if c_pengajuan is not None else "",
            )
            recs.append(rec)
        if recs:
            frames.append(pd.DataFrame(recs))
    summary = shopee_income_summary(wb)
    wb.close()
    if not frames:
        return pd.DataFrame(), seller, summary
    df = pd.concat(frames, ignore_index=True)
    df["created"] = pd.to_datetime(df["created"], errors="coerce")
    df["settled"] = pd.to_datetime(df["settled"], errors="coerce")
    df["line_key"] = (
        df["order_id"] + "|" + df["pengajuan"] + "|" + df["settled"].astype(str) + "|" + df["neto"].round(0).astype(str)
    )
    return df, seller, summary


def _parse_shopee_income_old(ws):
    """Format Laporan Penghasilan lama (≤ pertengahan 2026): sheet 'Income', 1 baris = 1 order, header di baris 'No. Pesanan'."""
    rows = list(ws.iter_rows(values_only=True))
    hi = _find_header_row(rows[:12], ["No. Pesanan", "Total Penghasilan"])
    if hi is None:
        return pd.DataFrame()
    header = [_s(c) for c in rows[hi]]
    idx = {h: i for i, h in enumerate(header) if h}
    g = lambda r, name: to_num(r[idx[name]]) if name in idx else 0.0
    gs = lambda r, names: sum(to_num(r[idx[n]]) for n in names if n in idx)
    ONGKIR = ["Ongkir Dibayar Pembeli", "Diskon Ongkir Ditanggung Jasa Kirim", "Gratis Ongkir dari Shopee",
              "Ongkir yang Diteruskan oleh Shopee ke Jasa Kirim", "Ongkos Kirim Pengembalian Barang",
              "Kembali ke Biaya Pengiriman Pengirim", "Pengembalian Biaya Kirim"]
    VOUCHER = ["Diskon Produk dari Shopee", "Voucher dari Penjual", "Voucher disponsor oleh Penjual",
               "Voucher co-fund disponsor oleh Penjual", "Cashback Koin dari Penjual", "Cashback Koin disponsori Penjual",
               "Cashback Koin Co-fund disponsori Penjual"]
    LAINNYA = ["Premi", "Biaya Program Hemat Biaya Kirim", "Biaya Transaksi", "Biaya Isi Saldo Otomatis (dari Penghasilan)"]
    recs = []
    for r in rows[hi + 1:]:
        oid = _s(r[idx["No. Pesanan"]])
        if not oid or oid.lower().startswith("total"):
            continue
        recs.append(dict(
            platform=SHOPEE, channel=SHOPEE, order_id=oid,
            created=_s(r[idx["Waktu Pesanan Dibuat"]]) if "Waktu Pesanan Dibuat" in idx else "",
            settled=_s(r[idx["Tanggal Dana Dilepaskan"]]) if "Tanggal Dana Dilepaskan" in idx else "",
            neto=g(r, "Total Penghasilan"),
            harga_produk=g(r, "Harga Asli Produk") + g(r, "Total Diskon Produk"),   # = harga setelah diskon penjual
            refund=g(r, "Jumlah Pengembalian Dana ke Pembeli"),
            ongkir_net=gs(r, ONGKIR),
            voucher_penjual=gs(r, VOUCHER),
            biaya_admin=g(r, "Biaya Administrasi"),
            biaya_proses=g(r, "Biaya Proses Pesanan"),
            biaya_xtra=g(r, "Biaya Layanan"),            # di format lama 'Biaya Layanan' = biaya layanan gratis ongkir XTRA
            biaya_layanan=0.0,
            komisi_ams=g(r, "Biaya Komisi AMS"),
            promosi_lain=g(r, "Biaya Kampanye"),
            biaya_lainnya=gs(r, LAINNYA),
            pph22=g(r, "Bea Masuk, PPN & PPh"),
            pengajuan=_s(r[idx["No. Pengajuan"]]) if "No. Pengajuan" in idx else "",
        ))
    return pd.DataFrame(recs)


# ---------------------------------------------------------------- TikTok orders

def parse_tiktok_orders(path):
    import io as _io
    if path.lower().endswith(".csv"):
        with open(path, "rb") as fh:
            raw = fh.read()
        txt = raw.replace(b"\x00", b"").decode("utf-8-sig", "ignore")     # file sebagian rusak: NUL dibuang, baris utuh tetap kebaca
        df = pd.read_csv(_io.StringIO(txt), dtype=str, keep_default_na=False, on_bad_lines="skip", engine="python")
    else:
        # pandas/openpyxl read_only salah baca dimensi file export TikTok → load penuh
        wb = openpyxl.load_workbook(path, data_only=True)
        rows = list(wb.worksheets[0].iter_rows(values_only=True))
        wb.close()
        hdr = [_s(c) for c in rows[0]]
        df = pd.DataFrame([["" if c is None else str(c) for c in r] + [""] * (len(hdr) - len(r)) for r in rows[1:]], columns=hdr)
        df = df.loc[:, [c for c in df.columns if c]]
    df.columns = [str(c).strip() for c in df.columns]
    for c in df.columns:
        df[c] = df[c].astype(str).str.strip()
    # buang baris deskripsi (export xlsx baris 2) & baris rusak/terpotong
    df = df[df["Order ID"].str.replace("\t", "").str.fullmatch(r"\d{6,}")].reset_index(drop=True)
    for c in ("Purchase Channel", "Seller SKU", "Variation", "Payment Method", "Cancel Reason", "Cancelation/Return Type", "SKU Seller Discount", "SKU Subtotal Before Discount"):
        if c not in df.columns:
            df[c] = ""
    out = pd.DataFrame()
    out["platform"] = [TIKTOK] * len(df)
    chan = df.get("Purchase Channel", pd.Series([""] * len(df)))
    out["channel"] = chan.where(chan.str.lower().eq("tokopedia"), TIKTOK).replace({"tokopedia": TOKOPEDIA})
    out.loc[chan.str.lower().eq("tokopedia"), "channel"] = TOKOPEDIA
    out["order_id"] = df["Order ID"]
    out["created"] = pd.to_datetime(df["Created Time"], dayfirst=True, errors="coerce")
    out["status_raw"] = df["Order Status"]
    out["status"] = out["status_raw"].replace({"": "(kosong)"})
    low = out["status_raw"].str.lower()
    out["is_batal"] = low.isin(["dibatalkan", "cancelled", "canceled"])
    out["is_unpaid"] = low.isin(["belum dibayar", "unpaid"])
    out["is_return"] = df.get("Cancelation/Return Type", pd.Series([""] * len(df))).str.lower().str.contains("return")
    out["sku"] = df["Seller SKU"].where(df["Seller SKU"] != "", df["SKU ID"])
    out["sku_id"] = df["SKU ID"]
    out["product"] = df["Product Name"]
    out["variation"] = df.get("Variation", "")
    out["qty"] = df["Quantity"].map(to_num)
    before = df["SKU Subtotal Before Discount"].map(to_num)
    seller_disc = df["SKU Seller Discount"].map(to_num)
    out["bruto"] = before - seller_disc                       # = after discount + diskon platform (ditanggung TikTok)
    out["harga_awal"] = before
    out["payment"] = df.get("Payment Method", "")
    out["cancel_reason"] = df.get("Cancel Reason", "")
    # identitas pembeli: TikTok menyamarkan username ("D***i a***iarini") dan telepon ("(+62)812*****60") →
    # kunci = gabungan keduanya (+ nama penerima) supaya pelanggan yang sama tetap terkumpul
    uname = df["Buyer Username"] if "Buyer Username" in df.columns else pd.Series([""] * len(df))
    rcpt = df["Recipient"] if "Recipient" in df.columns else pd.Series([""] * len(df))
    phone = df["Phone #"] if "Phone #" in df.columns else pd.Series([""] * len(df))
    out["buyer"] = uname.where(uname != "", rcpt)
    out["buyer_key"] = _buyer_key(uname, rcpt, phone)
    out["line_key"] = out["order_id"] + "|" + out["sku_id"] + "|" + out["variation"]
    return out


# ---------------------------------------------------------------- TikTok income

_TT_MAP = {
    "order_id": ["ID Pesanan/Penyesuaian", "Order/adjustment ID"],
    "type": ["Jenis transaksi", "Type"],
    "created": ["Waktu pemesanan", "Order created time"],
    "settled": ["Waktu pembayaran pesanan", "Order settled time"],
    "neto": ["Jumlah penyelesaian pembayaran", "Total settlement amount"],
    "revenue": ["Total Pendapatan", "Total Revenue"],
    "harga_produk": ["Subtotal setelah diskon penjual", "Subtotal after seller discounts"],
    "refund": ["Subtotal pengembalian dana setelah diskon penjual", "Refund subtotal after seller discounts"],
    "total_fees": ["Total Biaya", "Total Fees", "Total fees"],
    "komisi_platform": ["Biaya komisi platform", "Platform commission fee", "TikTok Shop commission fee"],
    "komisi_dinamis": ["Komisi dinamis", "Dynamic commission", "Dynamic Commission"],
    "biaya_proses": ["Biaya pemrosesan pesanan", "Order processing fee", "Flat fee"],
    "ongkir_net": ["Ongkir", "Shipping cost"],
    "pph22": ["PPh Pasal 22 dipungut", "Article 22 Income Tax withheld"],
    "adj": ["Jumlah penyesuaian", "Ajustment amount", "Adjustment amount"],
    "related": ["ID pesanan terkait", "Related order ID"],
    "source": ["Sumber pesanan", "Order Source"],
}
_TT_AFILIASI = [
    "Komisi Afiliasi", "Komisi mitra afiliasi", "Komisi Iklan Toko afiliasi", "Komisi iklan toko Mitra Afiliasi",
    "Deposit komisi afiliasi", "Pengembalian dana komisi afiliasi",
    "Affiliate Commission", "Affiliate commission", "Affiliate partner commission", "Affiliate Shop Ads commission",
    "Affiliate Partner shop ads commission", "TAP Shop Ads commission",
]


def parse_tiktok_income(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    summary = tiktok_income_summary(wb)
    wb.close()
    header = [_s(c) for c in rows[0]]
    idx = {}
    for i, h in enumerate(header):
        if h and h.lower() not in idx:
            idx[h.lower()] = i

    def col(key):
        for name in _TT_MAP[key]:
            if name.lower() in idx:
                return idx[name.lower()]
        return None

    cols = {k: col(k) for k in _TT_MAP}
    SUMMED = ("total_fees", "komisi_platform", "komisi_dinamis", "biaya_proses", "ongkir_net", "pph22")
    multi = {k: sorted({idx[n.lower()] for n in _TT_MAP[k] if n.lower() in idx}) for k in SUMMED}
    c_af = sorted({idx[n.lower()] for n in _TT_AFILIASI if n.lower() in idx})
    recs, adjs = [], []
    for r in rows[1:]:
        if r is None or all(c is None for c in r):
            continue
        g = lambda k: (_s(r[cols[k]]) if cols[k] is not None else "")
        v = lambda k: (sum(to_num(r[i]) for i in multi[k]) if k in multi else (to_num(r[cols[k]]) if cols[k] is not None else 0.0))
        typ = g("type")
        oid = g("order_id").strip("\t ")
        if typ.lower() in ("pesanan", "order"):
            af = sum(to_num(r[i]) for i in c_af)
            tf = v("total_fees")
            known = v("komisi_platform") + v("komisi_dinamis") + v("biaya_proses") + af + v("ongkir_net") + v("pph22")
            src = g("source").lower()
            recs.append(dict(
                platform=TIKTOK,
                channel=TOKOPEDIA if "tokopedia" in src else TIKTOK,
                order_id=oid, created=g("created"), settled=g("settled"),
                neto=v("neto"), harga_produk=v("harga_produk"), refund=v("refund"),
                komisi_platform=v("komisi_platform"), komisi_dinamis=v("komisi_dinamis"),
                biaya_proses=v("biaya_proses"), afiliasi=af, ongkir_net=v("ongkir_net"),
                pph22=v("pph22"), biaya_lainnya=tf - known, total_fees=tf,
            ))
        else:
            adjs.append(dict(platform=TIKTOK, type=typ, id=oid, related=g("related").strip("\t "),
                             settled=g("settled"), amount=v("neto") if v("neto") else v("adj")))
    df = pd.DataFrame(recs)
    if len(df):
        df["created"] = pd.to_datetime(df["created"], errors="coerce")
        df["settled"] = pd.to_datetime(df["settled"], errors="coerce")
        df["line_key"] = df["order_id"] + "|" + df["settled"].astype(str) + "|" + df["neto"].round(0).astype(str)
    adj = pd.DataFrame(adjs)
    if len(adj):
        adj["settled"] = pd.to_datetime(adj["settled"], errors="coerce")
    return df, adj, summary


# ---------------------------------------------------------------- TikTok belum settle (estimasi resmi)

def parse_tiktok_unsettled(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    hi = _find_header_row(rows[:10], ["ID pesanan terkait"])
    if hi is None:
        return pd.DataFrame()
    header = [_s(c) for c in rows[hi]]
    idx = {h: i for i, h in enumerate(header) if h}
    def pick(*names):
        for n in names:
            for h, i in idx.items():
                if h.startswith(n):
                    return i
        return None
    c_oid = pick("ID pesanan terkait", "Related order ID")
    c_est = pick("Estimasi jumlah penyelesaian", "Estimated settlement amount")
    c_when = pick("Estimasi waktu penyelesaian", "Estimated settlement time")
    c_why = pick("Alasan penunggakan", "Reason")
    c_created = pick("Waktu pemesanan", "Order created time")
    c_sku = pick("ID SKU", "SKU ID")
    c_prod = pick("Nama produk", "Product name")
    recs = []
    for r in rows[hi + 1:]:
        oid = _s(r[c_oid]).strip("\t ") if c_oid is not None else ""
        if not oid or oid == "/":
            continue
        recs.append(dict(platform=TIKTOK, order_id=oid, est_neto=to_num(r[c_est]) if c_est is not None else 0.0,
                         est_when=_s(r[c_when]) if c_when is not None else "", reason=_s(r[c_why]) if c_why is not None else "",
                         created=_s(r[c_created]) if c_created is not None else "", sku=_s(r[c_sku]) if c_sku is not None else "",
                         product=_s(r[c_prod]) if c_prod is not None else ""))
    df = pd.DataFrame(recs)
    if len(df):
        df["created"] = pd.to_datetime(df["created"], errors="coerce")
        df["line_key"] = df["order_id"] + "|" + df["sku"]
    return df


# ---------------------------------------------------------------- ringkasan resmi (buat audit)

SHOPEE_SUMMARY_MAP = {           # label di sheet Summary → kolom hasil parse
    "Harga Asli Produk": "harga_produk",
    "Total Diskon Produk": "diskon_produk",
    "Jumlah Pengembalian Dana ke Pembeli": "refund",
    "Biaya Administrasi": "biaya_admin",
    "Biaya Proses Pesanan": "biaya_proses",
    "Biaya Gratis Ongkir XTRA": "biaya_xtra",
    "Biaya Komisi AMS": "komisi_ams",
    "PPh 22": "pph22",
    "3. Total yang Dilepas": "neto",
}
TIKTOK_SUMMARY_MAP = {
    "Jumlah penyelesaian pembayaran": "settlement_total", "Total settlement amount": "settlement_total",
    "Subtotal setelah diskon penjual": "harga_produk", "Subtotal after seller discounts": "harga_produk",
    "Subtotal pengembalian dana setelah diskon penjual": "refund", "Refund subtotal after seller discounts": "refund",
    "Total Biaya": "total_fees", "Total Fees": "total_fees", "Total fees": "total_fees",
    "Biaya komisi platform": "komisi_platform", "Platform commission fee": "komisi_platform", "TikTok Shop commission fee": "komisi_platform",
    "Komisi dinamis": "komisi_dinamis", "Dynamic commission": "komisi_dinamis", "Dynamic Commission": "komisi_dinamis",
    "Biaya pemrosesan pesanan": "biaya_proses", "Order processing fee": "biaya_proses", "Flat fee": "biaya_proses",
    "Ongkir": "ongkir_net", "Shipping cost": "ongkir_net",
    "PPh Pasal 22 dipungut": "pph22", "Article 22 Income Tax withheld": "pph22",
}


def shopee_income_summary(wb):
    out = {}
    if "Summary" not in wb.sheetnames:
        return out
    seen = set()
    for r in wb["Summary"].iter_rows(values_only=True):
        cells = [_s(c) for c in r if c is not None and _s(c) != ""]
        if len(cells) < 2:
            continue
        lab, val = cells[0], cells[-1]
        if lab in SHOPEE_SUMMARY_MAP and lab not in seen:
            seen.add(lab)
            out[SHOPEE_SUMMARY_MAP[lab]] = to_num(val)
    return out


_TT_SUMMARY_LOWER = {k.lower(): v for k, v in TIKTOK_SUMMARY_MAP.items()}


def tiktok_income_summary(wb):
    out = {}
    ws = None
    for n in ("Laporan", "Reports"):
        if n in wb.sheetnames:
            ws = wb[n]
    if ws is None:
        return out
    seen_lab = set()
    for r in ws.iter_rows(values_only=True):
        cells = [_s(c) for c in r if c is not None and _s(c) != ""]
        if len(cells) < 2:
            continue
        lab, val = cells[0], cells[-1]
        key = _TT_SUMMARY_LOWER.get(lab.lower())
        if key and lab.lower() not in seen_lab:
            seen_lab.add(lab.lower())
            out[key] = out.get(key, 0.0) + to_num(val)
    return out


# ---------------------------------------------------------------- entry point + cache

PARSER_VERSION = "11"

KINDS = {
    "shopee_orders": "Shopee · Order.all",
    "shopee_income": "Shopee · Laporan Penghasilan",
    "tiktok_orders": "TikTok/Tokopedia · Semua Pesanan",
    "tiktok_income": "TikTok/Tokopedia · Laporan Penghasilan",
    "tiktok_unsettled": "TikTok · Belum settle (estimasi resmi)",
    "corrupt": "⚠ FILE RUSAK (isi kosong/NUL) — export ulang dari seller center",
}


def parse_file(path, cache_dir=None):
    """→ dict(kind, orders, income, adj, meta). Hasil di-cache per (nama, mtime, size)."""
    st = os.stat(path)
    key = hashlib.md5(f"{PARSER_VERSION}|{os.path.basename(path)}|{st.st_mtime}|{st.st_size}".encode()).hexdigest()
    cpath = os.path.join(cache_dir, key + ".pkl") if cache_dir else None
    if cpath and os.path.exists(cpath):
        with open(cpath, "rb") as fh:
            return pickle.load(fh)
    kind = detect(path)
    res = dict(kind=kind, orders=pd.DataFrame(), income=pd.DataFrame(), adj=pd.DataFrame(), unsettled=pd.DataFrame(),
               meta=dict(file=os.path.basename(path), mtime=st.st_mtime, seller=""))
    if kind == "shopee_orders":
        res["orders"] = parse_shopee_orders(path)
    elif kind == "shopee_income":
        res["income"], res["meta"]["seller"], res["meta"]["summary"] = parse_shopee_income(path)
    elif kind == "tiktok_orders":
        res["orders"] = parse_tiktok_orders(path)
    elif kind == "tiktok_income":
        res["income"], res["adj"], res["meta"]["summary"] = parse_tiktok_income(path)
    elif kind == "tiktok_unsettled":
        res["unsettled"] = parse_tiktok_unsettled(path)
    for k in ("orders", "income", "adj", "unsettled"):
        if len(res[k]):
            res[k]["file"] = res["meta"]["file"]
            res[k]["file_mtime"] = st.st_mtime
    # periode isi file
    for k in ("orders", "income", "unsettled"):
        d = res[k]
        if len(d) and "created" in d:
            res["meta"][k + "_rows"] = int(len(d))
            res["meta"][k + "_from"] = str(d["created"].min())[:10]
            res["meta"][k + "_to"] = str(d["created"].max())[:10]
    if cpath:
        os.makedirs(cache_dir, exist_ok=True)
        with open(cpath, "wb") as fh:
            pickle.dump(res, fh)
    return res
